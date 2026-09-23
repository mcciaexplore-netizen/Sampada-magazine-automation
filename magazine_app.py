from __future__ import annotations

import calendar
import csv
import json
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

import pandas as pd
import pdfplumber

from main import (
    create_links_template,
    extract_articles,
    load_config,
    make_qr_codes,
    metadata_and_script,
    narrate,
    render_videos,
)


ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "config.json"


def safe_segment(value: str) -> str:
    source = Path(value.strip())
    stem = "".join(ch for ch in source.stem if ch.isalnum() or ch in " -_").strip(" .")
    suffix = source.suffix.lower() if source.suffix.lower() in {".pdf", ".xlsx", ".csv"} else ""
    cleaned = f"{stem}{suffix}"
    if not cleaned or cleaned in {".", ".."}:
        raise ValueError("Folder name is empty or invalid.")
    return cleaned


def monthly_paths(base: Path, year: int, month: int) -> dict[str, Path]:
    month_name = calendar.month_name[month]
    root = base.resolve() / "Sampada" / str(year) / f"{month:02d} - {month_name}"
    paths = {
        "root": root,
        "audio": root / "Audio",
        "videos": root / "Videos",
        "qr": root / "QR Codes",
        "captions": root / "Captions",
        "data": root / "Data",
    }
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    return paths


def detect_qr_pdf_pages(pdf_path: Path, resolution: int = 180) -> list[int]:
    try:
        import cv2
        import numpy as np
    except ImportError:
        return []
    detector = cv2.QRCodeDetector()
    found: list[int] = []
    with pdfplumber.open(pdf_path) as pdf:
        for index, page in enumerate(pdf.pages):
            rgb = np.array(page.to_image(resolution=resolution, antialias=False).original.convert("RGB"))
            gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
            ok, _points = detector.detectMulti(gray)
            if ok:
                found.append(index + 1)
    return found


def select_articles_containing_qr(manifest: Path, qr_pdf_pages: list[int], use_existing_fallback: bool = True) -> pd.DataFrame:
    frame = pd.read_csv(manifest, encoding="utf-8-sig")
    config = load_config(CONFIG_PATH) if CONFIG_PATH.exists() else {}
    selected_pages = set(config.get("selected_pages", []))
    if qr_pdf_pages:
        frame["selected"] = frame.apply(
            lambda row: "yes" if any(int(row.pdf_start_page) <= page <= int(row.pdf_end_page) for page in qr_pdf_pages) else "no",
            axis=1,
        )
        frame["selection_reason"] = frame.apply(
            lambda row: "QR detected in article pages" if row["selected"] == "yes" else "No QR detected",
            axis=1,
        )
    else:
        if selected_pages:
            frame["selected"] = frame["printed_start_page"].astype(int).map(lambda p: "yes" if p in selected_pages else "no")
            frame["selection_reason"] = frame["selected"].map({"yes": "Config / proof selection", "no": "Unselected"})
        else:
            if not use_existing_fallback:
                frame["selected"] = "no"
            frame["selection_reason"] = frame["selected"].map(
                {"yes": "Proof fallback / review required", "no": "No QR detected"}
            )
    frame.to_csv(manifest, index=False, encoding="utf-8-sig")
    return frame


def save_manifest(frame: pd.DataFrame, manifest: Path) -> None:
    required = {
        "id", "selected", "printed_start_page", "printed_end_page", "pdf_start_page",
        "pdf_end_page", "language", "title", "article_file",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Missing manifest columns: {', '.join(sorted(missing))}")
    frame.to_csv(manifest, index=False, encoding="utf-8-sig")


def build_excel(metadata_csv: Path, output_xlsx: Path, month: int, year: int) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    work_dir = metadata_csv.parent
    manifest_path = work_dir / "manifest.csv"
    links_path = work_dir / "youtube_links.csv"
    captions_dir = work_dir / "captions"

    with metadata_csv.open("r", encoding="utf-8-sig") as f:
        meta_rows = list(csv.DictReader(f))

    manifest_rows = []
    if manifest_path.exists():
        with manifest_path.open("r", encoding="utf-8-sig") as f:
            manifest_rows = list(csv.DictReader(f))

    manifest_map = {r["id"]: r for r in manifest_rows}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    HEADER_FILL = PatternFill(start_color="102A36", end_color="102A36", fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color="1A4759", end_color="1A4759", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFB", end_color="F8FAFB", fill_type="solid")
    WHITE_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    REGULAR_FONT = Font(name="Segoe UI", size=10, color="222222")
    THIN_BORDER = Border(
        left=Side(style="thin", color="E0E0E0"), right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"), bottom=Side(style="thin", color="E0E0E0")
    )

    # Sheet 1: Plan
    ws_plan = wb.create_sheet(title="YouTube Plan")
    ws_plan.views.sheetView[0].showGridLines = True
    ws_plan.merge_cells("A1:K1")
    ws_plan["A1"] = f"Sampada Magazine – {calendar.month_name[month]} {year} YouTube Plan"
    ws_plan["A1"].font = TITLE_FONT
    ws_plan["A1"].fill = HEADER_FILL
    ws_plan["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_plan.row_dimensions[1].height = 40

    headers = ["ID", "Title", "Language", "Printed Pages", "YouTube URL", "Keywords", "Description", "Video File", "Audio File", "Caption File (.srt)", "Captions Preview"]
    ws_plan.append([])
    ws_plan.append(headers)
    ws_plan.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, start=1):
        c = ws_plan.cell(row=3, column=col_idx)
        c.font = WHITE_FONT
        c.fill = SUBHEADER_FILL
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 3, 4) else "left", vertical="center")

    for r_idx, r in enumerate(meta_rows, start=4):
        m = manifest_map.get(r["id"], {})
        printed = f"{m.get('printed_start_page', '')} - {m.get('printed_end_page', '')}"
        
        cap_file = r.get("caption_file", "")
        cap_path = None
        if cap_file and (work_dir / cap_file).exists():
            cap_path = work_dir / cap_file
        else:
            candidates = [
                work_dir / "captions" / f"{Path(r['script_file']).stem}.srt",
                work_dir / f"captions/{r['id']}.srt",
            ]
            for cand in candidates:
                if cand.exists():
                    cap_path = cand
                    cap_file = str(cand.relative_to(work_dir))
                    break
        if not cap_file:
            cap_file = f"captions/{Path(r.get('script_file', 'script')).stem}.srt"
        cap_prev = "\n".join(cap_path.read_text(encoding="utf-8").splitlines()[:12]) if cap_path and cap_path.exists() else ""
        vals = [
            r["id"], r["title"], "Marathi" if r.get("language") == "mr" else "English",
            printed, r.get("youtube_url", ""), r.get("keywords", ""), r.get("description", ""),
            r.get("video_file", ""), r.get("audio_file", ""), cap_file, cap_prev
        ]
        ws_plan.append(vals)
        ws_plan.row_dimensions[r_idx].height = 65
        for col_idx in range(1, len(vals) + 1):
            c = ws_plan.cell(row=r_idx, column=col_idx)
            c.font = REGULAR_FONT
            c.border = THIN_BORDER
            if r_idx % 2 == 0:
                c.fill = ZEBRA_FILL
            if col_idx in (1, 3, 4):
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (7, 11):
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def copy_outputs(work_dir: Path, drive_paths: dict[str, Path]) -> None:
    mappings = (("audio", "audio"), ("videos", "videos"), ("qr_codes", "qr"), ("captions", "captions"))
    for local_name, drive_key in mappings:
        source = work_dir / local_name
        if not source.exists() or drive_key not in drive_paths:
            continue
        for file in source.iterdir():
            if file.is_file():
                shutil.copy2(file, drive_paths[drive_key] / file.name)


def status_summary(work_dir: Path) -> dict[str, int]:
    def count(name: str, pattern: str) -> int:
        path = work_dir / name
        return len(list(path.glob(pattern))) if path.exists() else 0
    manifest = work_dir / "manifest.csv"
    selected = 0
    if manifest.exists():
        frame = pd.read_csv(manifest, encoding="utf-8-sig")
        selected_mask = frame["selected"].astype(str).str.lower().eq("yes")
        if "id" in frame.columns:
            selected_mask &= frame["id"].astype(str).str.lower().ne("full")
        selected = int(selected_mask.sum())
    return {
        "selected": selected,
        "audio": count("audio", "*.mp3"),
        "videos": count("videos", "*.mp4"),
        "qr": count("qr_codes", "*.png"),
    }
