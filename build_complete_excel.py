import csv
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
work_dir = ROOT / "work"
excel_path = work_dir / "Sampada_September_2026_YouTube_Plan.xlsx"

# Load metadata and manifest
with open(work_dir / "youtube_metadata.csv", "r", encoding="utf-8-sig") as f:
    metadata_rows = list(csv.DictReader(f))

with open(work_dir / "manifest.csv", "r", encoding="utf-8-sig") as f:
    manifest_rows = list(csv.DictReader(f))

with open(work_dir / "youtube_links.csv", "r", encoding="utf-8-sig") as f:
    links_rows = {r["id"]: r.get("youtube_url", "") for r in csv.DictReader(f)}

# QR placement lookup
qr_placements = {
    "p005": {"qr_pdf_page": 5, "printed_page": 7, "placement": "PDF page 5 → printed page 7 — QR after A New Chapter, Built on 92 Years of Purpose"},
    "p008": {"qr_pdf_page": 12, "printed_page": 14, "placement": "PDF page 12 → printed page 14 — QR after Deepening India’s Chemical Manufacturing Base"},
    "p015": {"qr_pdf_page": 15, "printed_page": 17, "placement": "PDF page 15 → printed page 17 — QR after The Chemistry Behind Make in India"},
    "p018": {"qr_pdf_page": 21, "printed_page": 23, "placement": "PDF page 21 → printed page 23 — QR after Maharashtra’s Chemical Opportunity"},
    "p024": {"qr_pdf_page": 25, "printed_page": 27, "placement": "PDF page 25 → printed page 27 — QR after Beyond the Pharmacy of the World"},
    "p028": {"qr_pdf_page": 28, "printed_page": 30, "placement": "PDF page 28 → printed page 30 — QR after India’s China + 1 Moment"},
    "p031": {"qr_pdf_page": 32, "printed_page": 34, "placement": "PDF page 32 → printed page 34 — QR after More Than a Molecule"},
    "p035": {"qr_pdf_page": 37, "printed_page": 39, "placement": "PDF page 37 → printed page 39 — QR after India’s Chemical Opportunity Gets Bigger"},
    "p040": {"qr_pdf_page": 41, "printed_page": 43, "placement": "PDF page 41 → printed page 43 — QR after Engineering Safer Industrial Pumps"},
    "p047": {"qr_pdf_page": 50, "printed_page": 52, "placement": "PDF page 50 → printed page 52 — QR after एमएसएमई क्षेत्रात सुधारण्यासाठी सरकारचे महत्त्वाचे पाऊल"},
    "p053": {"qr_pdf_page": 52, "printed_page": 54, "placement": "PDF page 52 → printed page 54 — QR after प्रक्रिया केलेल्या खाद्यपदार्थांच्या निर्यातीतील विश्वासार्ह नाव – कांचन कुलकर्णी"},
    "p055": {"qr_pdf_page": 53, "printed_page": 55, "placement": "PDF page 53 → printed page 55 — QR after New Alloy Delivers 50% More Strength, 400% Higher Ductility"},
}

manifest_map = {r["id"]: r for r in manifest_rows}

wb = openpyxl.Workbook()
wb.remove(wb.active)  # Remove default sheet

# Color tokens
HEADER_FILL = PatternFill(start_color="102A36", end_color="102A36", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="1A4759", end_color="1A4759", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8FAFB", end_color="F8FAFB", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FFF8E7", end_color="FFF8E7", fill_type="solid")
WHITE_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="102A36")
REGULAR_FONT = Font(name="Segoe UI", size=10, color="222222")

THIN_BORDER = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)

# -------------------------------------------------------------
# Sheet 1: YouTube Publishing Plan
# -------------------------------------------------------------
ws_plan = wb.create_sheet(title="YouTube Publishing Plan")
ws_plan.views.sheetView[0].showGridLines = True

ws_plan.merge_cells("A1:L1")
ws_plan["A1"] = "Sampada Magazine – September 2026 YouTube Publishing Plan"
ws_plan["A1"].font = TITLE_FONT
ws_plan["A1"].fill = HEADER_FILL
ws_plan["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_plan.row_dimensions[1].height = 40

plan_headers = [
    "ID", "Article Title", "Language", "Printed Pages", "PDF Pages",
    "YouTube URL", "Keywords / Tags", "Description", "Video File", "Audio File", "Caption File (.srt)", "Captions Preview (SRT)"
]

ws_plan.append([])
ws_plan.append(plan_headers)
ws_plan.row_dimensions[3].height = 28

for col_idx, h in enumerate(plan_headers, start=1):
    cell = ws_plan.cell(row=3, column=col_idx)
    cell.font = WHITE_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = Alignment(horizontal="center" if col_idx in (1, 3, 4, 5) else "left", vertical="center", wrap_text=True)

for row_idx, r in enumerate(metadata_rows, start=4):
    m = manifest_map.get(r["id"], {})
    printed_pages = f"{m.get('printed_start_page', '')} - {m.get('printed_end_page', '')}"
    pdf_pages = f"{m.get('pdf_start_page', '')} - {m.get('pdf_end_page', '')}"
    yt_url = links_rows.get(r["id"], "")
    caption_rel = r.get("caption_file", f"captions/{Path(r['script_file']).stem}.srt")
    caption_path = work_dir / caption_rel
    caption_text = caption_path.read_text(encoding="utf-8") if caption_path.exists() else ""
    preview_srt = "\n".join(caption_text.splitlines()[:12]) if caption_text else ""

    values = [
        r["id"],
        r["title"],
        "Marathi" if r["language"] == "mr" else "English",
        printed_pages,
        pdf_pages,
        yt_url,
        r["keywords"],
        r["description"],
        r["video_file"],
        r["audio_file"],
        caption_rel,
        preview_srt,
    ]
    ws_plan.append(values)
    ws_plan.row_dimensions[row_idx].height = 70
    
    is_zebra = (row_idx % 2 == 0)
    for col_idx in range(1, len(values) + 1):
        cell = ws_plan.cell(row=row_idx, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if is_zebra:
            cell.fill = ZEBRA_FILL
        
        if col_idx in (1, 3, 4, 5):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in (8, 12):  # Description and SRT Preview
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# -------------------------------------------------------------
# Sheet 2: YouTube Captions (Full Timed SRT)
# -------------------------------------------------------------
ws_captions = wb.create_sheet(title="YouTube Captions (SRT)")
ws_captions.views.sheetView[0].showGridLines = True

ws_captions.merge_cells("A1:F1")
ws_captions["A1"] = "YouTube Video Captions & Subtitles (Timed SRT Content)"
ws_captions["A1"].font = TITLE_FONT
ws_captions["A1"].fill = HEADER_FILL
ws_captions["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_captions.row_dimensions[1].height = 40

cap_headers = [
    "#", "Article ID", "Article Title", "Language", "Caption File (.srt)", "Full Timed Captions (SRT Format)"
]
ws_captions.append([])
ws_captions.append(cap_headers)
ws_captions.row_dimensions[3].height = 28

for col_idx, h in enumerate(cap_headers, start=1):
    cell = ws_captions.cell(row=3, column=col_idx)
    cell.font = WHITE_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4) else "left", vertical="center")

for idx, r in enumerate(metadata_rows, start=1):
    row_num = idx + 3
    caption_rel = r.get("caption_file", f"captions/{Path(r['script_file']).stem}.srt")
    caption_path = work_dir / caption_rel
    caption_full = caption_path.read_text(encoding="utf-8") if caption_path.exists() else ""

    values = [
        idx,
        r["id"],
        r["title"],
        "Marathi" if r["language"] == "mr" else "English",
        caption_rel,
        caption_full,
    ]
    ws_captions.append(values)
    ws_captions.row_dimensions[row_num].height = 90
    
    is_zebra = (row_num % 2 == 0)
    for col_idx in range(1, len(values) + 1):
        cell = ws_captions.cell(row=row_num, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if is_zebra:
            cell.fill = ZEBRA_FILL
        if col_idx in (1, 2, 4):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == 6:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# -------------------------------------------------------------
# Sheet 3: Full Narration Transcripts
# -------------------------------------------------------------
ws_scripts = wb.create_sheet(title="Narration Transcripts")
ws_scripts.views.sheetView[0].showGridLines = True

ws_scripts.merge_cells("A1:F1")
ws_scripts["A1"] = "Spoken Narration Scripts & Transcripts"
ws_scripts["A1"].font = TITLE_FONT
ws_scripts["A1"].fill = HEADER_FILL
ws_scripts["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_scripts.row_dimensions[1].height = 40

scr_headers = [
    "#", "Article ID", "Article Title", "Language", "Script File", "Full Narration Script (Text)"
]
ws_scripts.append([])
ws_scripts.append(scr_headers)
ws_scripts.row_dimensions[3].height = 28

for col_idx, h in enumerate(scr_headers, start=1):
    cell = ws_scripts.cell(row=3, column=col_idx)
    cell.font = WHITE_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4) else "left", vertical="center")

for idx, r in enumerate(metadata_rows, start=1):
    row_num = idx + 3
    script_path = work_dir / r["script_file"]
    script_full = script_path.read_text(encoding="utf-8") if script_path.exists() else ""

    values = [
        idx,
        r["id"],
        r["title"],
        "Marathi" if r["language"] == "mr" else "English",
        r["script_file"],
        script_full,
    ]
    ws_scripts.append(values)
    ws_scripts.row_dimensions[row_num].height = 80
    
    is_zebra = (row_num % 2 == 0)
    for col_idx in range(1, len(values) + 1):
        cell = ws_scripts.cell(row=row_num, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if is_zebra:
            cell.fill = ZEBRA_FILL
        if col_idx in (1, 2, 4):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == 6:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# -------------------------------------------------------------
# Sheet 4: QR Code Placement & Links
# -------------------------------------------------------------
ws_qr = wb.create_sheet(title="QR Code Placements")
ws_qr.views.sheetView[0].showGridLines = True

ws_qr.merge_cells("A1:G1")
ws_qr["A1"] = "QR Code Placement Instructions & Verification – September 2026 Proof"
ws_qr["A1"].font = TITLE_FONT
ws_qr["A1"].fill = HEADER_FILL
ws_qr["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_qr.row_dimensions[1].height = 40

qr_headers = [
    "#", "Article ID", "Article Title", "Printed Page", "PDF Proof Page", "Exact Placement Instruction", "YouTube URL"
]
ws_qr.append([])
ws_qr.append(qr_headers)
ws_qr.row_dimensions[3].height = 28

for col_idx, h in enumerate(qr_headers, start=1):
    cell = ws_qr.cell(row=3, column=col_idx)
    cell.font = WHITE_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4, 5) else "left", vertical="center")

for idx, r in enumerate(metadata_rows, start=1):
    qr_info = qr_placements.get(r["id"], {})
    row_num = idx + 3
    values = [
        idx,
        r["id"],
        r["title"],
        qr_info.get("printed_page", ""),
        qr_info.get("qr_pdf_page", ""),
        qr_info.get("placement", "End of article"),
        links_rows.get(r["id"], ""),
    ]
    ws_qr.append(values)
    ws_qr.row_dimensions[row_num].height = 30
    
    is_zebra = (row_num % 2 == 0)
    for col_idx in range(1, len(values) + 1):
        cell = ws_qr.cell(row=row_num, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if is_zebra:
            cell.fill = ZEBRA_FILL
        if col_idx in (1, 2, 4, 5):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# -------------------------------------------------------------
# Sheet 5: Full Manifest & Articles Catalog
# -------------------------------------------------------------
ws_manifest = wb.create_sheet(title="All Magazine Articles")
ws_manifest.views.sheetView[0].showGridLines = True

ws_manifest.merge_cells("A1:H1")
ws_manifest["A1"] = "Complete Table of Contents & Article Extraction Status"
ws_manifest["A1"].font = TITLE_FONT
ws_manifest["A1"].fill = HEADER_FILL
ws_manifest["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_manifest.row_dimensions[1].height = 40

man_headers = [
    "ID", "Selected for YouTube", "Printed Start", "Printed End", "PDF Start", "PDF End", "Language", "Title"
]
ws_manifest.append([])
ws_manifest.append(man_headers)
ws_manifest.row_dimensions[3].height = 28

for col_idx, h in enumerate(man_headers, start=1):
    cell = ws_manifest.cell(row=3, column=col_idx)
    cell.font = WHITE_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = Alignment(horizontal="center" if col_idx != 8 else "left", vertical="center")

for idx, m in enumerate(manifest_rows, start=4):
    is_selected = m.get("selected", "").lower() in ("yes", "y", "true", "1")
    values = [
        m["id"],
        "YES" if is_selected else "NO",
        int(m["printed_start_page"]),
        int(m["printed_end_page"]),
        int(m["pdf_start_page"]),
        int(m["pdf_end_page"]),
        "Marathi" if m["language"] == "mr" else "English",
        m["title"],
    ]
    ws_manifest.append(values)
    ws_manifest.row_dimensions[idx].height = 24
    
    for col_idx in range(1, len(values) + 1):
        cell = ws_manifest.cell(row=idx, column=col_idx)
        cell.font = BOLD_FONT if is_selected else REGULAR_FONT
        cell.border = THIN_BORDER
        if is_selected:
            cell.fill = HIGHLIGHT_FILL
        elif idx % 2 == 0:
            cell.fill = ZEBRA_FILL
        
        if col_idx in (1, 2, 3, 4, 5, 6, 7):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Explicit column widths
ws_plan.column_dimensions["A"].width = 10
ws_plan.column_dimensions["B"].width = 38
ws_plan.column_dimensions["C"].width = 14
ws_plan.column_dimensions["D"].width = 16
ws_plan.column_dimensions["E"].width = 14
ws_plan.column_dimensions["F"].width = 32
ws_plan.column_dimensions["G"].width = 36
ws_plan.column_dimensions["H"].width = 50
ws_plan.column_dimensions["I"].width = 36
ws_plan.column_dimensions["J"].width = 36
ws_plan.column_dimensions["K"].width = 36
ws_plan.column_dimensions["L"].width = 55

ws_captions.column_dimensions["A"].width = 6
ws_captions.column_dimensions["B"].width = 12
ws_captions.column_dimensions["C"].width = 38
ws_captions.column_dimensions["D"].width = 14
ws_captions.column_dimensions["E"].width = 36
ws_captions.column_dimensions["F"].width = 85

ws_scripts.column_dimensions["A"].width = 6
ws_scripts.column_dimensions["B"].width = 12
ws_scripts.column_dimensions["C"].width = 38
ws_scripts.column_dimensions["D"].width = 14
ws_scripts.column_dimensions["E"].width = 36
ws_scripts.column_dimensions["F"].width = 85

ws_qr.column_dimensions["A"].width = 6
ws_qr.column_dimensions["B"].width = 12
ws_qr.column_dimensions["C"].width = 42
ws_qr.column_dimensions["D"].width = 16
ws_qr.column_dimensions["E"].width = 16
ws_qr.column_dimensions["F"].width = 58
ws_qr.column_dimensions["G"].width = 35

ws_manifest.column_dimensions["A"].width = 10
ws_manifest.column_dimensions["B"].width = 24
ws_manifest.column_dimensions["C"].width = 14
ws_manifest.column_dimensions["D"].width = 14
ws_manifest.column_dimensions["E"].width = 12
ws_manifest.column_dimensions["F"].width = 12
ws_manifest.column_dimensions["G"].width = 14
ws_manifest.column_dimensions["H"].width = 50

try:
    wb.save(excel_path)
    print(f"Excel workbook with YouTube captions updated successfully: {excel_path}")
except PermissionError:
    alt_path = work_dir / "Sampada_September_2026_YouTube_Plan_Latest.xlsx"
    wb.save(alt_path)
    print(f"Original file was open in Excel. Saved updated copy to: {alt_path}")
